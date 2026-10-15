#!/usr/bin/env python3
"""Validate the Halvern workbook against the written rules and build dashboard.html.

Usage: python3 -I process.py <workbook.xlsx> <template.html> <output.html>
Nothing is corrected, filled or interpolated. Failures are recorded as issues.
"""
import datetime as dt
import json
import math
import re
import sys

import openpyxl

TODAY = dt.date(2026, 10, 15)
TOL = 0.05
QOQ_LIMIT = 0.30
SUFFIX = re.compile(r"\b(ltd|limited|plc|inc|group|holdings|llp)\b\.?", re.I)
QHDR = re.compile(r"^Q([1-4]) (\d{4})$")
FORMULA_MOIC = re.compile(r"^=\(G(\d+)\+H\1\)/F\1$")

RULES = [
    ("P1", "Portfolio A Company", "Non-empty text, unique. Exact duplicates and near-duplicates (same name once Ltd/Limited/plc/Inc/Group/Holdings/LLP removed, same fund and sector) are blocking; both rows quarantined.", "Blocking"),
    ("P2", "Portfolio B Fund", "Must match a fund on Fund Summary.", "Blocking"),
    ("P3", "Portfolio C Sector, D Geography", "Non-empty text.", "Blocking"),
    ("P4", "Portfolio E Entry Date", "True date, not after 15 Oct 2026, not before fund vintage year. Text dates are not parsed. Drops the row from date views only.", "Warning"),
    ("P5", "Portfolio F Invested Capital", "Number above 0.", "Blocking"),
    ("P6", "Portfolio G Current Valuation", "Number of 0 or more. Blank is not treated as 0.", "Blocking"),
    ("P7", "Portfolio H Realised Proceeds", "Number of 0 or more. Blank is not treated as 0.", "Blocking"),
    ("P8", "Portfolio I MOIC", "Must be the formula (G+H)/F. A typed value, or one that differs from the recomputed value, is blocking.", "Blocking"),
    ("P9", "Portfolio J Status", "Exactly Active or Exited. Exited needs Valuation = 0 and Realised > 0. Active needs Valuation > 0.", "Blocking"),
    ("Q1", "Marks A Company", "Names must match the Portfolio tab one-to-one.", "Warning"),
    ("Q2", "Marks B:O cells", "Headers 'Qn YYYY' and contiguous. Cells numeric >= 0 or blank. Text such as '12.4 (est)' is flagged, not parsed.", "Warning"),
    ("Q3", "Marks B:O gaps", "Blank allowed only before the first mark. A later blank is a gap and is not interpolated.", "Warning"),
    ("Q4", "Marks first mark", "First mark quarter should equal entry quarter (or Q1 2023 if entry was earlier).", "Warning"),
    ("Q5", "Marks QoQ", "Move of more than 30% between consecutive valid marks is flagged.", "Warning"),
    ("Q6", "Marks vs Portfolio", "Latest mark must equal Current Valuation (Active) or Realised Proceeds (Exited), tolerance 0.05.", "Warning"),
    ("Q7", "Marks vs Portfolio", "First mark compared with Invested Capital (information only).", "Info"),
    ("F1", "Fund Summary", "Fund names unique, vintage integer, Net IRR numeric. Net IRR shown as reported and cannot be verified without cash flows.", "Info"),
    ("F2", "Fund Summary C:G", "SUMIF ranges must cover every Portfolio row.", "Warning"),
    ("F3", "Fund Summary C:G", "Recompute Invested, Value, MOIC, DPI from validated rows; bridge to workbook formula results (blanks as 0).", "Warning"),
    ("F4", "Fund Summary E,H", "TVPI formula must differ from Gross MOIC; flagged if identical.", "Warning"),
    ("F5", "Whole workbook", "Formulas with no stored results are flagged.", "Info"),
]

issues = []


def issue(sev, rule, sheet, cell, entity, found, treatment):
    issues.append(dict(sev=sev, rule=rule, sheet=sheet, cell=cell, entity=entity,
                       found=found, treatment=treatment))


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def norm(name):
    return re.sub(r"\s+", " ", SUFFIX.sub("", name or "")).strip().lower()


def main(path, tpl, out):
    wb = openpyxl.load_workbook(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    pf, qm, fs = wb["Portfolio"], wb["Quarterly Marks"], wb["Fund Summary"]

    # ---- Fund Summary (F1)
    funds = []
    for r in range(2, fs.max_row + 1):
        name = fs.cell(r, 1).value
        if name is None:
            continue
        vint, irr = fs.cell(r, 2).value, fs.cell(r, 6).value
        if not isinstance(vint, int):
            issue("Warning", "F1", "Fund Summary", f"B{r}", name, f"Vintage {vint!r} is not an integer", "Fund vintage not used")
        if not is_num(irr):
            issue("Warning", "F1", "Fund Summary", f"F{r}", name, f"Net IRR {irr!r} is not numeric", "Not shown")
        funds.append(dict(row=r, name=str(name).strip(), vintage=vint, irr=irr if is_num(irr) else None,
                          f_invested=fs.cell(r, 3).value, f_value=fs.cell(r, 4).value,
                          f_moic=fs.cell(r, 5).value, f_dpi=fs.cell(r, 7).value, f_tvpi=fs.cell(r, 8).value))
    names = [f["name"].lower() for f in funds]
    for f in funds:
        if names.count(f["name"].lower()) > 1:
            issue("Warning", "F1", "Fund Summary", f"A{f['row']}", f["name"], "Duplicate fund name", "Flagged")
    fund_map = {f["name"]: f for f in funds}
    for f in funds:
        issue("Info", "F1", "Fund Summary", f"F{f['row']}", f["name"],
              f"Net IRR {f['irr']}% is hard-coded; no cash flows in the workbook to recompute it", "Shown as reported, labelled unverified")

    # F2: SUMIF ranges cover all portfolio rows
    last_pf = max(r for r in range(2, pf.max_row + 1) if pf.cell(r, 1).value is not None)
    for f in funds:
        for col in (3, 4):
            m = re.findall(r"\$([A-Z])\$(\d+):\$([A-Z])\$(\d+)", str(fs.cell(f["row"], col).value))
            if not m or any(int(x[3]) < last_pf or int(x[1]) > 2 for x in m):
                issue("Warning", "F2", "Fund Summary", f"{'C' if col == 3 else 'D'}{f['row']}", f["name"],
                      f"SUMIF range does not span Portfolio rows 2-{last_pf}", "Workbook totals may omit rows")
    # F4: TVPI identical to Gross MOIC
    for f in funds:
        if f["f_tvpi"] and f["f_tvpi"] == f["f_moic"]:
            issue("Warning", "F4", "Fund Summary", f"H{f['row']}", f["name"],
                  f"TVPI formula {f['f_tvpi']} is identical to Gross MOIC; TVPI is normally net of fees and carry, so this is a gross figure under a net label",
                  "Shown as reported; not presented as a net measure")
    # F5: no stored results
    fs_v = wbv["Fund Summary"]
    if all(fs_v.cell(f["row"], 3).value is None for f in funds):
        issue("Info", "F5", "Whole workbook", "Fund Summary C:H, Portfolio I", "-",
              "Formula cells have no stored results (the file was never calculated in Excel)",
              "All formula results recomputed here from the stated formulas")

    # ---- Portfolio
    rows = []
    for r in range(2, last_pf + 1):
        g = lambda c: pf.cell(r, c).value
        rows.append(dict(row=r, company=g(1), fund=g(2), sector=g(3), geo=g(4), entry=g(5),
                         inv=g(6), val=g(7), real=g(8), moic_cell=g(9), status=g(10)))
    for p in rows:
        p["block"] = []
        p["entry_ok"] = False
        p["entry_q"] = None
        name = p["company"]
        n = p["row"]

        def block(rule, col, found):
            p["block"].append(rule)
            issue("Blocking", rule, "Portfolio", f"{col}{n}", name, found, "Row excluded from all headline figures")

        if not isinstance(name, str) or not name.strip():
            block("P1", "A", "Company name missing"); name = f"(row {n})"; p["company"] = name
        else:
            p["company"] = name = name.strip()
        if p["fund"] not in fund_map:
            block("P2", "B", f"Fund {p['fund']!r} not on Fund Summary")
        for c, key, lab in (("C", "sector", "Sector"), ("D", "geo", "Geography")):
            if not isinstance(p[key], str) or not p[key].strip():
                block("P3", c, f"{lab} missing")
        # P4
        e = p["entry"]
        vint = fund_map.get(p["fund"], {}).get("vintage")
        if isinstance(e, dt.datetime):
            d = e.date()
            if d > TODAY:
                issue("Warning", "P4", "Portfolio", f"E{n}", name, f"Entry date {d} is in the future", "Excluded from date views")
            elif isinstance(vint, int) and d.year < vint:
                issue("Warning", "P4", "Portfolio", f"E{n}", name, f"Entry date {d} is before fund vintage {vint}", "Excluded from date views")
            else:
                p["entry_ok"] = True
                p["entry_q"] = (d.year, (d.month - 1) // 3 + 1)
        else:
            issue("Warning", "P4", "Portfolio", f"E{n}", name,
                  f"Entry date is not a date value: {e!r}" if e is not None else "Entry date missing",
                  "Not parsed or corrected; excluded from date views only")
        # P5-P7
        if not is_num(p["inv"]):
            block("P5", "F", f"Invested Capital is {p['inv']!r}, not a number")
        elif p["inv"] <= 0:
            block("P5", "F", f"Invested Capital is {p['inv']}; must be above 0")
        if p["val"] is None:
            block("P6", "G", "Current Valuation is blank (not assumed to be 0)")
        elif not is_num(p["val"]) or p["val"] < 0:
            block("P6", "G", f"Current Valuation is {p['val']!r}; must be a number of 0 or more")
        if p["real"] is None:
            block("P7", "H", "Realised Proceeds is blank (not assumed to be 0)")
        elif not is_num(p["real"]) or p["real"] < 0:
            block("P7", "H", f"Realised Proceeds is {p['real']!r}; must be a number of 0 or more")
        # P9
        st = p["status"]
        if st not in ("Active", "Exited"):
            block("P9", "J", f"Status is {st!r}; must be Active or Exited")
        elif is_num(p["val"]) and is_num(p["real"]):
            if st == "Exited" and not (p["val"] == 0 and p["real"] > 0):
                block("P9", "J", f"Exited but Valuation = {p['val']} and Realised = {p['real']}")
            if st == "Active" and not p["val"] > 0:
                block("P9", "J", f"Active but Valuation = {p['val']}")
        # P8
        mc = p["moic_cell"]
        ok_inputs = is_num(p["inv"]) and p["inv"] > 0 and is_num(p["val"]) and is_num(p["real"])
        p["moic"] = (p["val"] + p["real"]) / p["inv"] if ok_inputs else None
        if isinstance(mc, str) and FORMULA_MOIC.match(mc) and int(FORMULA_MOIC.match(mc).group(1)) == n:
            pass
        elif is_num(mc):
            note = ""
            if p["moic"] is not None:
                note = f"; recomputed (G+H)/F = {p['moic']:.2f}x"
            block("P8", "I", f"MOIC is a typed value {mc}x, not the formula{note}. Either the MOIC or the inputs are wrong; cannot tell which")
        else:
            block("P8", "I", f"MOIC cell is {mc!r}; expected formula (G+H)/F")

    # P1 duplicates
    seen = {}
    for p in rows:
        seen.setdefault(p["company"].lower(), []).append(p)
    for k, grp in seen.items():
        if len(grp) > 1:
            for p in grp:
                if "P1" not in p["block"]:
                    p["block"].append("P1")
                issue("Blocking", "P1", "Portfolio", f"A{p['row']}", p["company"],
                      f"Exact duplicate company name (rows {', '.join(str(x['row']) for x in grp)})", "All duplicate rows excluded")
    groups = {}
    for p in rows:
        groups.setdefault((norm(p["company"]), p["fund"], p["sector"]), []).append(p)
    for grp in groups.values():
        if len(grp) > 1 and len({x["company"].lower() for x in grp}) > 1:
            desc = "; ".join(f"row {x['row']} '{x['company']}' invested {x['inv']}, valued {x['val']}, entry {x['entry'].date() if isinstance(x['entry'], dt.datetime) else x['entry']}" for x in grp)
            for x in grp:
                if "P1" not in x["block"]:
                    x["block"].append("P1")
                issue("Blocking", "P1", "Portfolio", f"A{x['row']}", x["company"],
                      f"Possible duplicate: same company, fund and sector under two names. {desc}", "Both rows excluded; cannot tell which is correct, or whether both are real")

    # ---- Quarterly Marks
    hdr = [qm.cell(1, c).value for c in range(2, qm.max_column + 1)]
    quarters, qkeys = [], []
    for h in hdr:
        m = QHDR.match(str(h).strip()) if h is not None else None
        if m:
            quarters.append(str(h).strip()); qkeys.append((int(m.group(2)), int(m.group(1))))
        else:
            issue("Warning", "Q2", "Quarterly Marks", "row 1", "-", f"Header {h!r} is not 'Qn YYYY'", "Column ignored")
    for a, b in zip(qkeys, qkeys[1:]):
        if (b[0] * 4 + b[1]) - (a[0] * 4 + a[1]) != 1:
            issue("Warning", "Q2", "Quarterly Marks", "row 1", "-", f"Quarter headers not contiguous at {a} -> {b}", "Flagged")
    first_q = qkeys[0]
    pmap = {p["company"].lower(): p for p in rows}
    marks = {}
    qm_names = set()
    for r in range(2, qm.max_row + 1):
        nm = qm.cell(r, 1).value
        if nm is None:
            continue
        nm = str(nm).strip()
        qm_names.add(nm.lower())
        p = pmap.get(nm.lower())
        if p is None:
            issue("Warning", "Q1", "Quarterly Marks", f"A{r}", nm, "Company not on Portfolio tab", "Ignored")
            continue
        vals = []
        for i in range(len(quarters)):
            v = qm.cell(r, 2 + i).value
            cell = f"{openpyxl.utils.get_column_letter(2 + i)}{r}"
            if v is None:
                vals.append(None)
            elif is_num(v) and v >= 0:
                vals.append(float(v))
            else:
                vals.append(None)
                issue("Warning", "Q2", "Quarterly Marks", cell, nm, f"{quarters[i]}: {v!r} is not a valid number", "Excluded from the series (shown as a gap)")
        # raw first mark (a malformed cell counts as present for Q3/Q4 purposes)
        raw = [qm.cell(r, 2 + i).value for i in range(len(quarters))]
        first_i = next((i for i, v in enumerate(raw) if v is not None), None)
        if first_i is None:
            issue("Warning", "Q3", "Quarterly Marks", f"A{r}", nm, "No marks at all", "Flagged")
        else:
            for i in range(first_i, len(quarters)):
                if raw[i] is None:
                    issue("Warning", "Q3", "Quarterly Marks", f"{openpyxl.utils.get_column_letter(2 + i)}{r}", nm,
                          f"{quarters[i]}: blank after first mark ({quarters[first_i]})", "Gap kept; not interpolated; trend point omitted")
            # Q4
            if p["entry_q"]:
                exp = max(p["entry_q"], first_q)
                if qkeys[first_i] != exp:
                    issue("Warning", "Q4", "Quarterly Marks", f"A{r}", nm,
                          f"First mark is {quarters[first_i]} but entry date implies Q{exp[1]} {exp[0]}", "Flagged")
            # Q7 info
            fv = vals[first_i]
            if fv is not None and is_num(p["inv"]) and abs(fv - p["inv"]) > TOL:
                issue("Info", "Q7", "Quarterly Marks", f"{openpyxl.utils.get_column_letter(2 + first_i)}{r}", nm,
                      f"First mark {fv} differs from Invested Capital {p['inv']}", "Information only")
        # Q5
        prev = None
        for i, v in enumerate(vals):
            if v is None:
                continue
            if prev is not None and qkeys[i][0] * 4 + qkeys[i][1] - (qkeys[prev[0]][0] * 4 + qkeys[prev[0]][1]) == 1 and prev[1] > 0:
                ch = v / prev[1] - 1
                if abs(ch) > QOQ_LIMIT:
                    issue("Warning", "Q5", "Quarterly Marks", f"{openpyxl.utils.get_column_letter(2 + i)}{r}", nm,
                          f"{quarters[i]} mark moves {ch:+.0%} on the prior quarter", "Flagged")
            prev = (i, v)
        # Q6
        last_i = max((i for i, v in enumerate(vals) if v is not None), default=None)
        ref = None
        if p["status"] == "Active":
            ref, lab = p["val"], "Current Valuation"
        elif p["status"] == "Exited":
            ref, lab = p["real"], "Realised Proceeds"
        if last_i is not None and is_num(ref):
            if last_i != len(quarters) - 1:
                issue("Warning", "Q6", "Quarterly Marks", f"A{r}", nm, f"Latest valid mark is {quarters[last_i]}, not {quarters[-1]}", "Compared at latest valid mark")
            if abs(vals[last_i] - ref) > TOL:
                issue("Warning", "Q6", "Quarterly Marks", f"{openpyxl.utils.get_column_letter(2 + last_i)}{r}", nm,
                      f"Latest mark {vals[last_i]} vs Portfolio {lab} {ref}: difference {vals[last_i] - ref:+.1f}", "Shown on dashboard; Portfolio figure used for headlines")
        marks[p["company"]] = vals
        p["marks"] = vals
    for p in rows:
        if p["company"].lower() not in qm_names:
            issue("Warning", "Q1", "Portfolio", f"A{p['row']}", p["company"], "No row on Quarterly Marks tab", "No trend shown")
            p["marks"] = [None] * len(quarters)

    # ---- Validated figures
    good = [p for p in rows if not p["block"]]

    def agg(ps):
        inv = sum(p["inv"] for p in ps)
        val = sum(p["val"] for p in ps)
        real = sum(p["real"] for p in ps)
        return dict(n=len(ps), inv=inv, val=val, real=real, tv=val + real,
                    moic=(val + real) / inv if inv else None, dpi=real / inv if inv else None)

    def xl_sum(ps, k):
        return sum(p[k] for p in ps if is_num(p[k]))  # Excel SUMIF treats blanks as 0

    fund_out = []
    for f in funds:
        allf = [p for p in rows if p["fund"] == f["name"]]
        gf = [p for p in good if p["fund"] == f["name"]]
        v = agg(gf) if gf else None
        wi, wv = xl_sum(allf, "inv"), xl_sum(allf, "val") + xl_sum(allf, "real")
        wr = xl_sum(allf, "real")
        wk = dict(inv=wi, tv=wv, moic=wv / wi if wi else None, dpi=wr / wi if wi else None)
        fund_out.append(dict(name=f["name"], vintage=f["vintage"], irr=f["irr"], validated=v, workbook=wk,
                             rows_total=len(allf), rows_valid=len(gf)))
        if v and wk["inv"] is not None and (abs(v["inv"] - wk["inv"]) > TOL or abs(v["tv"] - wk["tv"]) > TOL):
            issue("Warning", "F3", "Fund Summary", f"C{f['row']}:G{f['row']}", f["name"],
                  f"Workbook formulas give invested {wk['inv']:.1f}, value {wk['tv']:.1f}, MOIC {wk['moic']:.2f}x over all {len(allf)} rows; "
                  f"validated rows ({len(gf)}) give invested {v['inv']:.1f}, value {v['tv']:.1f}, MOIC {v['moic']:.2f}x",
                  "Dashboard shows validated figures with the workbook figures beside them")

    # trend: per fund, sum of valid marks; point shown only if every expected company has a clean mark
    trend = []
    for f in funds:
        gf = [p for p in good if p["fund"] == f["name"]]
        pts = []
        for i, q in enumerate(quarters):
            exp = [p for p in gf if any(v is not None for v in p["marks"][: i + 1]) or False]
            exp = [p for p in gf if next((j for j, v in enumerate(p["marks"]) if v is not None), 99) <= i]
            have = [p for p in exp if p["marks"][i] is not None]
            complete = bool(exp) and len(have) == len(exp)
            pts.append(dict(q=q, v=sum(p["marks"][i] for p in have) if complete else None,
                            n=len(have), expected=len(exp)))
        trend.append(dict(fund=f["name"], pts=pts))

    def mix(key):
        d = {}
        for p in good:
            d[p[key]] = d.get(p[key], 0) + p["inv"]
        return sorted(({"k": k, "v": v} for k, v in d.items()), key=lambda x: -x["v"])

    comp_issues = {}
    for i in issues:
        comp_issues.setdefault(i["entity"], []).append(i["rule"])

    companies = []
    for p in rows:
        companies.append(dict(
            company=p["company"], fund=p["fund"], sector=p["sector"], geo=p["geo"], status=p["status"],
            entry=p["entry"].date().isoformat() if isinstance(p["entry"], dt.datetime) else None,
            inv=p["inv"] if is_num(p["inv"]) else None, val=p["val"] if is_num(p["val"]) else None,
            real=p["real"] if is_num(p["real"]) else None,
            moic=p["moic"] if not p["block"] else None,
            excluded=bool(p["block"]), block=sorted(set(p["block"])),
            rules=sorted(set(comp_issues.get(p["company"], []))), row=p["row"],
            last_mark=next((v for v in reversed(p["marks"]) if v is not None), None)))

    order = {"Blocking": 0, "Warning": 1, "Info": 2}
    issues.sort(key=lambda i: (order[i["sev"]], i["rule"], i["sheet"], i["cell"]))
    sev_count = {s: sum(1 for i in issues if i["sev"] == s) for s in order}

    data = dict(
        generated=TODAY.isoformat(), source=path.split("/")[-1],
        rows_total=len(rows), rows_valid=len(good), sev=sev_count,
        headline=agg(good) if good else None,
        workbook_all=dict(inv=xl_sum(rows, "inv"), tv=xl_sum(rows, "val") + xl_sum(rows, "real"),
                          moic=(xl_sum(rows, "val") + xl_sum(rows, "real")) / xl_sum(rows, "inv"),
                          dpi=xl_sum(rows, "real") / xl_sum(rows, "inv")),
        funds=fund_out, sector=mix("sector"), geo=mix("geo"), companies=companies,
        quarters=quarters, trend=trend, issues=issues,
        rules=[dict(id=a, where=b, rule=c, sev=d) for a, b, c, d in RULES],
        excluded=[dict(company=c["company"], rules=c["block"]) for c in companies if c["excluded"]],
    )
    html = open(tpl, encoding="utf-8").read().replace("__DATA__", json.dumps(data, ensure_ascii=False))
    open(out, "w", encoding="utf-8").write(html)
    print(json.dumps(dict(valid=len(good), total=len(rows), sev=sev_count, headline=data["headline"],
                          workbook_all=data["workbook_all"]), indent=1))


if __name__ == "__main__":
    main(*sys.argv[1:4])
